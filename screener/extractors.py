"""把 PDF / Word / PowerPoint 等檔案轉成純文字。"""
import shutil
import subprocess
import tempfile
from pathlib import Path

SUPPORTED_EXTS = {".pdf", ".docx", ".doc", ".pptx", ".ppt", ".txt", ".md", ".rtf", ".odt"}

# PDF 抽出的文字少於這個字數，視為掃描檔／圖片型 PDF
MIN_PDF_CHARS = 200
# 其他格式少於這個字數視為空檔
MIN_TEXT_CHARS = 20


class ExtractionError(Exception):
    pass


def extract_text(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".pdf":
        return _pdf(path)
    if ext == ".docx":
        return _docx(path)
    if ext == ".pptx":
        return _pptx(path)
    if ext in {".txt", ".md"}:
        return path.read_text(encoding="utf-8", errors="ignore")
    if ext in {".doc", ".ppt", ".rtf", ".odt"}:
        return _via_libreoffice(path)
    raise ExtractionError(f"不支援的檔案格式：{ext}")


def _pdf(path: Path) -> str:
    import pdfplumber

    parts = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            parts.append(page.extract_text() or "")
            for table in page.extract_tables() or []:
                for row in table:
                    parts.append(" | ".join(c or "" for c in row))
    return "\n".join(parts)


def _docx(path: Path) -> str:
    import docx

    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            cells = []
            for c in row.cells:
                if c.text not in cells:  # 合併儲存格會重複出現
                    cells.append(c.text)
            parts.append(" | ".join(cells))
    for section in d.sections:
        parts.extend(p.text for p in section.header.paragraphs)
    return "\n".join(p for p in parts if p.strip())


def _pptx(path: Path) -> str:
    from pptx import Presentation

    prs = Presentation(str(path))
    parts = []
    for i, slide in enumerate(prs.slides, 1):
        parts.append(f"--- 第 {i} 頁 ---")
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
        if slide.has_notes_slide:
            parts.append(slide.notes_slide.notes_text_frame.text)
    return "\n".join(p for p in parts if p.strip())


def _via_libreoffice(path: Path) -> str:
    """舊版 .doc / .ppt 透過 LibreOffice 轉成新格式再讀。"""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        for candidate in (
            r"C:\Program Files\LibreOffice\program\soffice.exe",
            "/Applications/LibreOffice.app/Contents/MacOS/soffice",
        ):
            if Path(candidate).exists():
                soffice = candidate
                break
    if not soffice:
        raise ExtractionError(
            f"{path.suffix} 為舊版格式，需安裝 LibreOffice 才能讀取；或請另存為 .docx / .pptx / .pdf"
        )
    target = {".ppt": "pptx"}.get(path.suffix.lower(), "docx")
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run(
            [soffice, "--headless", "--convert-to", target, "--outdir", tmp, str(path)],
            check=True, capture_output=True, timeout=120,
        )
        converted = Path(tmp) / (path.stem + "." + target)
        if not converted.exists():
            raise ExtractionError("LibreOffice 轉檔失敗")
        return _pptx(converted) if target == "pptx" else _docx(converted)
